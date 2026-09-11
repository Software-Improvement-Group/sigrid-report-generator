#  Copyright Software Improvement Group
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.

import pytest
from pptx import Presentation
from pptx.util import Inches

from report_generator.generator.domain.portfolio.maintainability_portfolio.statistics import (
    maintainability_portfolio_stats,
)
from report_generator.generator.placeholders.implementations.text.maintainability_portfolio import (
    portfolio_maint_avg_delta_param,
)
from report_generator.generator.placeholders.rendering import pptx as render
from report_generator.generator.utils.constants import MaintMetric


@pytest.fixture
def prime_metric_averages():
    """Prime the portfolio stats singleton's caches and clean them up afterwards."""

    def _prime(metric: MaintMetric, start_average, end_average):
        maintainability_portfolio_stats.__dict__["statistics"] = {
            "metric-averages": {
                metric.to_json_name(): {
                    "start-average": start_average,
                    "end-average": end_average,
                }
            }
        }

    yield _prime

    maintainability_portfolio_stats.__dict__.pop("statistics", None)


def _presentation_with_text(text):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(2))
    textbox.text_frame.paragraphs[0].text = text
    return presentation


def _resolve(presentation, metric: MaintMetric):
    placeholder = portfolio_maint_avg_delta_param
    key = placeholder.key.format(parameter=metric)
    value_cb = lambda: placeholder.value(metric)  # noqa: E731
    placeholder.resolve_pptx(presentation, key, value_cb)
    return key


def _only_run(presentation):
    paragraph = presentation.slides[0].shapes[0].text_frame.paragraphs[0]
    return paragraph.runs[0]


def test_key_is_derived_from_custom_key_template():
    assert (
        portfolio_maint_avg_delta_param.key == "PORTFOLIO_MAINT_AVG_DELTA_{parameter}"
    )


def test_increase_renders_signed_value_in_green(prime_metric_averages):
    prime_metric_averages(MaintMetric.UNIT_COMPLEXITY, 3.2, 3.5)
    key = portfolio_maint_avg_delta_param.key.format(
        parameter=MaintMetric.UNIT_COMPLEXITY
    )
    presentation = _presentation_with_text(key)

    _resolve(presentation, MaintMetric.UNIT_COMPLEXITY)

    run = _only_run(presentation)
    assert run.text == "+0.30"
    assert run.font.color.rgb == render.FIVE_STAR_COLOR


def test_decrease_renders_signed_value_in_red(prime_metric_averages):
    prime_metric_averages(MaintMetric.UNIT_SIZE, 4.0, 3.99)
    key = portfolio_maint_avg_delta_param.key.format(parameter=MaintMetric.UNIT_SIZE)
    presentation = _presentation_with_text(key)

    _resolve(presentation, MaintMetric.UNIT_SIZE)

    run = _only_run(presentation)
    assert run.text == "-0.01"
    assert run.font.color.rgb == render.ONE_STAR_COLOR


def test_unchanged_renders_equals_in_blue(prime_metric_averages):
    prime_metric_averages(MaintMetric.DUPLICATION, 3.5, 3.5)
    key = portfolio_maint_avg_delta_param.key.format(parameter=MaintMetric.DUPLICATION)
    presentation = _presentation_with_text(key)

    _resolve(presentation, MaintMetric.DUPLICATION)

    run = _only_run(presentation)
    assert run.text == "="
    assert run.font.color.rgb == render.SIG_BLUE_COLOR


def test_no_data_at_either_boundary_renders_equals_in_blue(prime_metric_averages):
    """A submetric absent from every snapshot must render as unchanged rather than a spurious
    large delta derived from the near-zero sentinel used elsewhere for empty weighted averages."""
    prime_metric_averages(MaintMetric.MODULE_COUPLING, None, None)
    key = portfolio_maint_avg_delta_param.key.format(
        parameter=MaintMetric.MODULE_COUPLING
    )
    presentation = _presentation_with_text(key)

    _resolve(presentation, MaintMetric.MODULE_COUPLING)

    run = _only_run(presentation)
    assert run.text == "="
    assert run.font.color.rgb == render.SIG_BLUE_COLOR
