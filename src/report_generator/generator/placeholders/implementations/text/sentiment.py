#  Copyright Software Improvement Group
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.

"""Text placeholders whose wording and color both follow from a single number."""

from abc import ABC
from collections.abc import Callable
from typing import NamedTuple

from docx.document import Document
from pptx.presentation import Presentation

from report_generator.generator.placeholders import rendering
from report_generator.generator.placeholders.formatting import formatters
from report_generator.generator.placeholders.implementations.base import (
    MultiParameterList,
    ParameterizedPlaceholder,
    ParameterList,
    Placeholder,
    function_name_to_placeholder_key,
)
from report_generator.generator.placeholders.rendering.common import (
    FontColor,
    FontProperties,
)
from report_generator.generator.utils.constants.sentiment import Sentiment

from .base import AbstractTextPlaceholder, as_multi_parameter_list


class SentimentText(NamedTuple):
    """How a number becomes colored text: its wording and its sentiment. Pass one of the presets
    below to `sentiment_text_placeholder`."""

    format: Callable[[float], str]
    sentiment: Callable[[float], Sentiment]


#: Renders a numeric delta as a signed value (e.g. +0.01, -0.01, =), green for an increase, red
#: for a decrease and blue when unchanged.
SIGNED_DELTA = SentimentText(
    format=formatters.format_signed_delta, sentiment=formatters.delta_sentiment
)

#: Renders a star rating as its position relative to market average: 'below' red (< 2.5),
#: 'average' blue (2.5 - 3.4) and 'above' green (>= 3.5).
MARKET_AVERAGE = SentimentText(
    format=formatters.format_market_average,
    sentiment=formatters.market_average_sentiment,
)


class _AbstractSentimentTextPlaceholder(AbstractTextPlaceholder, ABC):
    """Renders a number as text colored by the sentiment of that number. Subclasses supply the
    wording and sentiment rules through `__sentiment_text__`, and `value()` returns the raw
    number so that both the wording and the color can be derived from it."""

    __sentiment_text__: SentimentText

    @classmethod
    def resolve_pptx(
        cls, presentation: Presentation, key: str, value_cb: Callable[[], float]
    ) -> None:
        paragraphs = rendering.pptx.find_text_in_presentation(presentation, key)
        if not paragraphs:
            return
        number = value_cb()
        font = FontProperties(
            color=FontColor(
                rgb=rendering.pptx.sentiment_color(
                    cls.__sentiment_text__.sentiment(number)
                )
            )
        )
        rendering.pptx.update_many_paragraphs(
            paragraphs, key, cls.__sentiment_text__.format(number), font
        )

    @classmethod
    def resolve_docx(
        cls, document: Document, key: str, value_cb: Callable[[], float]
    ) -> None:
        cls._resolve_with_adapter(
            cls._DOCX_ADAPTER,
            document,
            key,
            lambda: cls.__sentiment_text__.format(value_cb()),
        )


def sentiment_text_placeholder(
    sentiment_text: SentimentText, custom_key: str | None = None
) -> Callable[[Callable[[], float]], type[Placeholder]]:
    """Turn a function returning a number into a text placeholder that renders the number as the
    wording of `sentiment_text`, colored by its sentiment. Coloring is applied in PowerPoint;
    Word renders the wording without color."""

    def decorator(number_func: Callable[[], float]) -> type[Placeholder]:
        class SentimentTextPlaceholder(_AbstractSentimentTextPlaceholder):
            __doc__ = number_func.__doc__ if number_func.__doc__ else None
            __sentiment_text__ = sentiment_text
            key = (
                custom_key
                if custom_key
                else function_name_to_placeholder_key(number_func.__name__)
            )

            @classmethod
            def value(cls) -> float:
                return number_func()

        return SentimentTextPlaceholder

    return decorator


def parameterized_sentiment_text_placeholder(
    sentiment_text: SentimentText,
    custom_key: str,
    parameters: ParameterList | MultiParameterList,
) -> Callable[[Callable[..., float]], type[ParameterizedPlaceholder]]:
    """Turn a function returning a number into a parameterized text placeholder that expands per
    parameter, rendering each number as the wording of `sentiment_text`, colored by its
    sentiment. Coloring is applied in PowerPoint; Word renders the wording without color."""
    normalized_parameters = as_multi_parameter_list(parameters)

    def decorator(number_func: Callable[..., float]) -> type[ParameterizedPlaceholder]:
        class ParameterizedSentimentTextPlaceholder(
            ParameterizedPlaceholder, _AbstractSentimentTextPlaceholder
        ):
            __doc__ = number_func.__doc__ if number_func.__doc__ else None
            __sentiment_text__ = sentiment_text
            key = custom_key
            allowed_parameters = normalized_parameters

            @classmethod
            def value(cls, *args) -> float:
                return number_func(*args)

        return ParameterizedSentimentTextPlaceholder

    return decorator
